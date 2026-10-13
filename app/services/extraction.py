"""Text extraction from CV files (PDF via pdfplumber, DOCX via python-docx).

Besides the plain text, extraction records layout signals that the format
checks use: text that sits in tables, text boxes or page headers/footers,
which ATS parsers often skip or scramble. For DOCX these are read from the
XML and are reliable; for PDF they are heuristics.
"""

import io
import math
import re
import zipfile
from dataclasses import dataclass, field
from typing import Literal

FileKind = Literal["pdf", "docx"]

# Below this many letters we treat the file as scanned / image-only.
MIN_LETTERS_FOR_TEXT = 100
# Rough words-per-page for a dense one-page CV, used when DOCX has no page count.
WORDS_PER_PAGE_ESTIMATE = 600

_FINAL_HEBREW_LETTERS = "םןץףך"
_HEBREW_WORD_RE = re.compile(r"[א-ת]")
_MIRRORED = str.maketrans("()[]{}<>", ")(][}{><")

_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"


class UnsupportedFileError(ValueError):
    pass


class ExtractionError(ValueError):
    pass


@dataclass
class Extraction:
    file_kind: FileKind
    text: str
    page_count: int | None
    page_count_estimated: bool = False
    table_text: str = ""
    textbox_text: str = ""
    header_footer_text: str = ""
    # Heuristic layout observations, e.g. "multi_column", "pdf_tables".
    layout_notes: list[str] = field(default_factory=list)
    hebrew_order_fixed: bool = False

    @property
    def letter_count(self) -> int:
        return sum(1 for char in self.text if char.isalpha())

    @property
    def text_ok(self) -> bool:
        return self.letter_count >= MIN_LETTERS_FOR_TEXT


def detect_kind(filename: str, data: bytes) -> FileKind:
    """Decide the file type from its content, not only its name."""
    if data[:5] == b"%PDF-":
        return "pdf"
    if data[:4] == b"PK\x03\x04":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if "word/document.xml" in archive.namelist():
                    return "docx"
        except zipfile.BadZipFile:
            pass
    lower = filename.lower()
    if lower.endswith(".doc"):
        raise UnsupportedFileError("Old .doc files are not supported. Save it as .docx or PDF.")
    raise UnsupportedFileError("Only PDF and DOCX files are supported.")


def extract(data: bytes, kind: FileKind) -> Extraction:
    if kind == "pdf":
        return _extract_pdf(data)
    return _extract_docx(data)


# --- PDF -------------------------------------------------------------------


def _extract_pdf(data: bytes) -> Extraction:
    import pdfplumber  # imported here: it is slow to import and only needed on upload

    page_texts: list[str] = []
    edge_words: list[str] = []
    split_lines = 0
    multiword_lines = 0
    table_count = 0
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            page_count = len(pdf.pages)
            for page in pdf.pages:
                page_texts.append(page.extract_text() or "")
                words = page.extract_words()
                page_split, page_multi = _count_split_lines(words, float(page.width))
                split_lines += page_split
                multiword_lines += page_multi
                edge_words.extend(_edge_words(words, float(page.height)))
                for table in page.find_tables():
                    rows = table.extract()
                    if len(rows) >= 2 and max((len(row) for row in rows), default=0) >= 2:
                        table_count += 1
    except Exception as exc:  # pdfplumber raises several unrelated exception types
        raise ExtractionError(f"Could not read the PDF: {exc}") from exc

    text = "\n".join(page_texts)
    result = Extraction(file_kind="pdf", text=text, page_count=page_count)

    if looks_like_reversed_hebrew(text):
        result.text = fix_reversed_hebrew(text)
        result.hebrew_order_fixed = True

    # Many lines with a wide horizontal gap in the middle suggest two columns.
    if split_lines >= 8 and multiword_lines and split_lines / multiword_lines >= 0.25:
        result.layout_notes.append("multi_column")
    if table_count:
        result.layout_notes.append("pdf_tables")
    result.header_footer_text = " ".join(edge_words)
    return result


def _count_split_lines(words: list[dict], page_width: float) -> tuple[int, int]:
    """Count lines that have a wide gap between two words (a column gutter)."""
    lines: dict[int, list[dict]] = {}
    for word in words:
        lines.setdefault(round(float(word["top"]) / 3), []).append(word)
    split = 0
    multi = 0
    for line_words in lines.values():
        if len(line_words) < 2:
            continue
        multi += 1
        line_words.sort(key=lambda w: float(w["x0"]))
        gaps = [
            float(current["x0"]) - float(previous["x1"])
            for previous, current in zip(line_words, line_words[1:], strict=False)
        ]
        if max(gaps) > page_width * 0.12:
            split += 1
    return split, multi


def _edge_words(words: list[dict], page_height: float) -> list[str]:
    """Words in the top or bottom 6% of the page, where headers/footers live."""
    top_limit = page_height * 0.06
    bottom_limit = page_height * 0.94
    return [
        str(word["text"])
        for word in words
        if float(word["top"]) < top_limit or float(word["bottom"]) > bottom_limit
    ]


def looks_like_reversed_hebrew(text: str) -> bool:
    """Detect Hebrew that came out of a PDF in visual (reversed) order.

    Final letter forms (ם ן ץ ף ך) only appear at the end of a word. If many
    Hebrew words start with one, the characters are reversed.
    """
    starts = 0
    ends = 0
    for token in text.split():
        letters = "".join(char for char in token if "א" <= char <= "ת")
        if len(letters) < 2:
            continue
        if letters[0] in _FINAL_HEBREW_LETTERS:
            starts += 1
        if letters[-1] in _FINAL_HEBREW_LETTERS:
            ends += 1
    return starts >= 3 and starts > ends


def fix_reversed_hebrew(text: str) -> str:
    """Restore logical order for runs of Hebrew words in each line.

    Within a run of consecutive Hebrew tokens, the token order is reversed and
    each token's characters are reversed. Latin tokens are left alone. This is
    an approximation of the bidi algorithm that is good enough for keyword
    matching; numbers inside Hebrew runs may still end up out of place.
    """
    fixed_lines: list[str] = []
    for line in text.split("\n"):
        tokens = line.split(" ")
        output: list[str] = []
        run: list[str] = []
        for token in tokens:
            if _HEBREW_WORD_RE.search(token):
                run.append(token[::-1].translate(_MIRRORED))
                continue
            if run:
                output.extend(reversed(run))
                run = []
            output.append(token)
        if run:
            output.extend(reversed(run))
        fixed_lines.append(" ".join(output))
    return "\n".join(fixed_lines)


# --- DOCX ------------------------------------------------------------------


def _extract_docx(data: bytes) -> Extraction:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError(f"Could not read the DOCX: {exc}") from exc

    body_parts: list[str] = []
    table_parts: list[str] = []
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            body_parts.append(Paragraph(child, document).text)
        elif child.tag == qn("w:tbl"):
            table_text = _docx_table_text(Table(child, document))
            body_parts.append(table_text)
            table_parts.append(table_text)

    textbox_text = "\n".join(_docx_textbox_paragraphs(document.element.body))

    header_footer_parts: list[str] = []
    seen_parts: set[int] = set()
    for section in document.sections:
        for part in (
            section.header, section.footer,
            section.first_page_header, section.first_page_footer,
            section.even_page_header, section.even_page_footer,
        ):
            if part.is_linked_to_previous:
                continue
            element_id = id(part._element)
            if element_id in seen_parts:
                continue
            seen_parts.add(element_id)
            texts = [paragraph.text for paragraph in part.paragraphs]
            texts.extend(_docx_table_text(table) for table in part.tables)
            texts.extend(_docx_textbox_paragraphs(part._element))
            header_footer_parts.extend(text for text in texts if text.strip())
    header_footer_text = "\n".join(header_footer_parts)

    # Matching should see everything, including content ATS parsers may skip.
    full_text = "\n".join(
        part for part in (header_footer_text, "\n".join(body_parts), textbox_text) if part.strip()
    )

    page_count, estimated = _docx_page_count(data, full_text)
    return Extraction(
        file_kind="docx",
        text=full_text,
        page_count=page_count,
        page_count_estimated=estimated,
        table_text="\n".join(table_parts),
        textbox_text=textbox_text,
        header_footer_text=header_footer_text,
    )


def _docx_table_text(table) -> str:  # type: ignore[no-untyped-def]
    rows: list[str] = []
    for row in table.rows:
        cells: list[str] = []
        previous = None
        for cell in row.cells:
            # Merged cells are returned once per grid column; skip repeats.
            if previous is not None and cell._tc is previous:
                continue
            previous = cell._tc
            if cell.text.strip():
                cells.append(cell.text.strip())
        if cells:
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def _docx_textbox_paragraphs(element) -> list[str]:  # type: ignore[no-untyped-def]
    """Paragraph texts inside text boxes, skipping the duplicate VML fallback copies."""
    paragraphs: list[str] = []
    for textbox in element.iter(f"{{{_W_NS}}}txbxContent"):
        if any(ancestor.tag == _MC_FALLBACK for ancestor in textbox.iterancestors()):
            continue
        for paragraph in textbox.iter(f"{{{_W_NS}}}p"):
            text = "".join(node.text or "" for node in paragraph.iter(f"{{{_W_NS}}}t"))
            if text.strip():
                paragraphs.append(text)
    return paragraphs


def _docx_page_count(data: bytes, text: str) -> tuple[int | None, bool]:
    """Page count from docProps/app.xml (written by Word on save), else an estimate."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if "docProps/app.xml" in archive.namelist():
                xml = archive.read("docProps/app.xml").decode("utf-8", errors="ignore")
                match = re.search(r"<Pages>(\d+)</Pages>", xml)
                if match and int(match.group(1)) > 0:
                    return int(match.group(1)), False
    except zipfile.BadZipFile:
        pass
    words = len(text.split())
    if words == 0:
        return None, True
    return max(1, math.ceil(words / WORDS_PER_PAGE_ESTIMATE)), True
