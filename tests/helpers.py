"""Builders for test inputs: sample JDs, DOCX and PDF files."""

import io
from pathlib import Path

FIXTURES = Path(__file__).parent / "fixtures"


def read_fixture(name: str) -> str:
    return (FIXTURES / name).read_text(encoding="utf-8")


def make_docx(
    paragraphs: list[str],
    table_rows: list[list[str]] | None = None,
    header_text: str | None = None,
    textbox_text: str | None = None,
    pages: int | None = None,
) -> bytes:
    """Build a DOCX in memory with optional table, header and text box."""
    from docx import Document
    from docx.oxml import parse_xml

    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for row_index, row in enumerate(table_rows):
            for col_index, value in enumerate(row):
                table.cell(row_index, col_index).text = value
    if header_text:
        document.sections[0].header.is_linked_to_previous = False
        document.sections[0].header.paragraphs[0].text = header_text
    if textbox_text:
        # A minimal VML text box, as older Word versions write them.
        xml = (
            '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'xmlns:v="urn:schemas-microsoft-com:vml"><w:r><w:pict><v:shape><v:textbox>'
            f'<w:txbxContent><w:p><w:r><w:t>{textbox_text}</w:t></w:r></w:p></w:txbxContent>'
            '</v:textbox></v:shape></w:pict></w:r></w:p>'
        )
        document.element.body.append(parse_xml(xml))
    buffer = io.BytesIO()
    document.save(buffer)
    data = buffer.getvalue()
    if pages is not None:
        data = _with_page_count(data, pages)
    return data


def _with_page_count(data: bytes, pages: int) -> bytes:
    """Write docProps/app.xml with a <Pages> value, as Word does on save."""
    import zipfile

    source = zipfile.ZipFile(io.BytesIO(data))
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "docProps/app.xml":
                content = (
                    '<?xml version="1.0" encoding="UTF-8"?><Properties '
                    'xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties">'
                    f"<Pages>{pages}</Pages></Properties>"
                ).encode()
            target.writestr(item, content)
    return output.getvalue()


def make_pdf(lines: list[str], two_columns: bool = False, blank: bool = False) -> bytes:
    """Build a simple PDF with reportlab. blank=True gives a page with no text (like a scan)."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    width, height = A4
    if blank:
        pdf.rect(50, 50, 200, 200, fill=1)
    else:
        y = height - 72
        for line in lines:
            if two_columns:
                pdf.drawString(50, y, line[:30])
                pdf.drawString(width / 2 + 40, y, line[:30])
            else:
                pdf.drawString(50, y, line)
            y -= 16
    pdf.showPage()
    pdf.save()
    return buffer.getvalue()


SAMPLE_CV_LINES = [
    "Niv Example",
    "niv@example.com | 054-123-4567 | github.com/niv-example",
    "Skills: C, C++, Python, Linux, Wireshark, Scapy, TCP/IP, Git, Docker",
    "Network traffic analyzer in Python, Jan 2025 - Present",
    "Checksum utilities in C, 2024",
    "B.Sc. Computer Science, Open University of Israel, 2021 - 2026",
    "Languages: Hebrew (native), English (fluent)",
    "Built packet parsing tools and signature based detection for pcap files.",
    "Wrote unit tests with pytest and set up CI with GitHub Actions.",
]
